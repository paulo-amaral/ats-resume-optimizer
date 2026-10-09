"""Deterministic checks: truthfulness invariants, keyword coverage, de-slop, ATS file hygiene."""
from __future__ import annotations

import re
from pathlib import Path

from .doc import Doc
from .i18n import fold

STANDARD_SECTIONS = ["summary", "experience", "education", "skills", "certifications"]


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _has(term: str, low: str) -> bool:
    return re.search(rf"(?<!\w){re.escape(term)}(?!\w)", low) is not None


# ---------------------------------------------------------------------------
# validate: truthfulness invariants (resume JSON vs source JSON)
# ---------------------------------------------------------------------------

def validate(resume: dict, source: dict | None = None) -> dict:
    failures: list[str] = []
    warnings: list[str] = []

    for key in ("basics", "experience"):
        if key not in resume:
            failures.append(f"missing required top-level section: '{key}'")
    if not resume.get("basics", {}).get("name"):
        failures.append("basics.name is missing")

    if source:
        roles_by_company: dict[str, list[dict]] = {}
        for e in source.get("experience", []):
            roles_by_company.setdefault(_norm(e.get("company", "")), []).append(e)

        for exp in resume.get("experience", []):
            company = _norm(exp.get("company", ""))
            if company and company not in roles_by_company:
                failures.append(f"FABRICATED employer not in source: '{exp.get('company')}'")
                continue
            matches = roles_by_company.get(company, [])
            same_title = [m for m in matches if _norm(m.get("title", "")) == _norm(exp.get("title", ""))]
            if matches and not same_title:
                warnings.append(
                    f"title '{exp.get('title')}' at '{exp.get('company')}' not found verbatim in source "
                    "(confirm it is a legitimate rewording, not a promotion inflation)"
                )
            # Compare dates against the role with the same title, so two roles at one employer don't clash.
            candidates = same_title or matches
            start = _norm(exp.get("start", ""))
            if start and candidates and all(
                _norm(m.get("start", "")) and _norm(m.get("start", "")) != start for m in candidates
            ):
                warnings.append(
                    f"start date differs from source for '{exp.get('company')}' "
                    f"({exp.get('start')} vs {', '.join(m.get('start', '') for m in candidates)})"
                )

        sb, rb = source.get("basics", {}), resume.get("basics", {})
        for f in ("name", "email", "phone"):
            if sb.get(f) and rb.get(f) and _norm(sb[f]) != _norm(rb[f]):
                failures.append(f"identity field '{f}' changed from source ('{sb[f]}' -> '{rb[f]}')")

        for sec in ("experience", "education", "skills", "certifications"):
            if source.get(sec) and not resume.get(sec):
                warnings.append(f"section '{sec}' populated in source but absent in resume (confirm intentional)")

    return {"invariants_passed": not failures, "failures": failures, "warnings": warnings}


# ---------------------------------------------------------------------------
# coverage: labeled keyword-coverage ESTIMATE
# ---------------------------------------------------------------------------

STOPWORDS = set(
    """a an the and or but if then else for to of in on at by with from into over under
    is are was were be been being have has had do does did will would shall should can
    could may might must this that these those it its as we you your our their they them
    he she his her i me my mine us who whom which what when where why how not no yes all
    any some each every both few more most other such own same so than too very just also
    about above after again against because before below between during down out off up
    while here there once only per via etc
    de da do das dos e ou em no na nos nas um uma uns umas para por pelo pela pelos pelas
    com sem que se ao aos à às como mais menos muito seu sua seus suas o os as é são ser
    estar ter ser será sobre entre até desde também já não sim este esta estes estas esse
    essa isso isto nosso nossa nossos nossas você vocês ele ela eles elas""".split()
)

GENERIC_JD = set(
    """experience years work working role roles responsibilities responsible ability able
    strong excellent good great team teams member members company candidate candidates
    join looking seeking required requirement requirements preferred plus bonus nice
    qualifications qualification skills skill knowledge understanding proficiency
    proficient familiarity familiar including include includes job position
    opportunity opportunities environment fast paced growth grow help support
    ensure ensuring build building develop developing manage managing lead leading
    provide providing deliver delivering across within using use used new best world
    class self starter driven oriented detail passionate motivated communication
    communicate collaborate collaboration stakeholders stakeholder cross functional
    day daily week weekly month monthly year annual please apply benefits salary
    location remote hybrid onsite full time part contract senior junior mid level
    domain exposure hands-on handson end
    experiência experiências anos ano trabalho trabalhar vaga vagas empresa candidato
    candidata candidatos requisitos requisito obrigatório obrigatórios desejável desejáveis
    diferencial diferenciais conhecimento conhecimentos habilidade habilidades capacidade
    atuação atuar atividades atividade responsabilidades responsável benefícios salário
    local remoto remota híbrido híbrida presencial integral júnior junior pleno sênior
    nível área time equipe equipes buscamos procuramos oportunidade principais
    perfil boa bom forte sólida sólido vivência desenvolver desenvolvimento apoiar
    garantir participar realizar suporte""".split()
)

PREFERRED_MARKERS = (
    "preferred", "nice to have", "nice-to-have", "bonus", "a plus", "plus:", "desirable", "ideally",
    "diferencial", "diferenciais", "desejável", "desejáveis", "será um plus", "é um plus",
)
REQUIRED_MARKERS = (
    "required", "requirements", "must have", "must-have", "qualifications", "what you'll need",
    "you have", "you bring", "requisitos", "obrigatório", "obrigatórios", "imprescindível",
    "o que buscamos", "o que esperamos", "pré-requisitos",
)

# Unicode letters so 'gestão' and 'análise' stay one token; keeps c++, c#, node.js intact.
_TOKEN = re.compile(r"[^\W\d_][\w+#.\-]*[\w+#]|[^\W\d_]")


def tokenize(text: str) -> list[str]:
    return _TOKEN.findall(text.lower())


def stem(t: str) -> str:
    """Naive plural strip so 'pipelines' matches 'pipeline'. Leaves c++, c# alone."""
    t = fold(t)
    if len(t) > 4 and t.endswith("ies"):
        return t[:-3] + "y"
    if len(t) > 3 and t.endswith("s") and not t.endswith("ss"):
        return t[:-1]
    return t


def is_content(t: str) -> bool:
    if t in STOPWORDS or t in GENERIC_JD:
        return False
    return len(t) >= 3 or any(c.isdigit() or c in "+#" for c in t)


def extract_keywords(jd_text: str) -> tuple[set[str], set[str]]:
    """Return (required, preferred) unigrams + adjacent bigrams, tiered by section markers.

    Bigrams only form inside a punctuation-delimited phrase, so 'Python, SQL' never yields 'python sql'.
    """
    tier = "required"
    req: set[str] = set()
    pref: set[str] = set()
    for line in jd_text.splitlines():
        low = line.lower()
        if any(m in low for m in PREFERRED_MARKERS):
            tier = "preferred"
        elif any(m in low for m in REQUIRED_MARKERS):
            tier = "required"
        bucket = pref if tier == "preferred" else req
        for phrase in re.split(r"[,/();:•]|\s[–—]\s|\b(?:and|or|with|e|ou|com)\b", low):
            toks = tokenize(phrase)
            bucket.update(t for t in toks if is_content(t))
            for a, b in zip(toks, toks[1:], strict=False):
                if is_content(a) and is_content(b):
                    bucket.add(f"{a} {b}")
    return req, pref - req


def coverage_of(keywords: set[str], text: str) -> tuple[set[str], set[str]]:
    """Unigram hit = stemmed token present. Bigram hit = phrase present OR both tokens present."""
    folded = fold(text)
    rtokens = {stem(t) for t in tokenize(text)}
    present, missing = set(), set()
    for kw in keywords:
        if " " in kw:
            a, b = kw.split(" ", 1)
            hit = fold(kw) in folded or (stem(a) in rtokens and stem(b) in rtokens)
        else:
            hit = stem(kw) in rtokens
        (present if hit else missing).add(kw)
    return present, missing


def coverage(doc: Doc, jd_text: str) -> dict:
    req, pref = extract_keywords(jd_text)
    req_present, req_missing = coverage_of(req, doc.text)
    pref_present, pref_missing = coverage_of(pref, doc.text)
    req_cov = len(req_present) / len(req) if req else 1.0
    pref_cov = len(pref_present) / len(pref) if pref else 1.0
    quant = sum(bool(re.search(r"\d", b)) for b in doc.bullets) / len(doc.bullets) if doc.bullets else 0.0
    completeness = sum(1 for s in STANDARD_SECTIONS if doc.sections.get(s)) / len(STANDARD_SECTIONS)

    content = {k: v for k, v in doc.sections.items() if k in ("summary", "experience", "skills")}
    all_kw = req | pref
    with_kw = sum(1 for v in content.values() if coverage_of(all_kw, v)[0])
    distribution = with_kw / len(content) if content else 0.0

    score = 0.40 * req_cov + 0.20 * pref_cov + 0.20 * quant + 0.10 * completeness + 0.10 * distribution
    return {
        "_note": "ESTIMATE of keyword fit + quality. NOT a real ATS internal score.",
        "estimated_score_pct": round(score * 100, 1),
        "band": "strong" if score >= 0.85 else "moderate" if score >= 0.70 else "weak",
        "components": {
            "required_coverage_pct": round(req_cov * 100, 1),
            "preferred_coverage_pct": round(pref_cov * 100, 1),
            "quantification_rate_pct": round(quant * 100, 1),
            "section_completeness_pct": round(completeness * 100, 1),
            "keyword_distribution_pct": round(distribution * 100, 1),
        },
        "missing_required_keywords": sorted(req_missing),
        "missing_preferred_keywords": sorted(pref_missing),
    }


# ---------------------------------------------------------------------------
# deslop: AI-tells, filler, inflated verbs (EN + PT-BR)
# ---------------------------------------------------------------------------

INFLATED = {
    "spearheaded": "led", "orchestrated": "coordinated", "architected": "designed",
    "leveraged": "used", "utilized": "used", "facilitated": "ran/helped",
    "championed": "advocated for", "pioneered": "introduced", "harnessed": "used",
    "helmed": "led", "navigated": "managed",
    "alavanquei": "aumentei / usei", "alavancou": "aumentou / usou", "alavancar": "aumentar / usar",
    "capitaneei": "liderei", "orquestrei": "coordenei", "potencializei": "aumentei (com número)",
    "fomentei": "incentivei / criei",
}
BUZZWORDS = {
    "synergy", "synergize", "paradigm", "robust", "scalable", "actionable", "impactful",
    "best-in-class", "cutting-edge", "world-class", "next-level", "results-driven",
    "detail-oriented", "team player", "hard worker", "go-getter", "self-starter",
    "dynamic", "passionate", "holistic", "seamless", "game-changer", "value-add",
    "mission-critical", "deep dive", "move the needle", "low-hanging fruit",
    "think outside the box",
    "sinergia", "proativo", "proativa", "dinâmico", "dinâmica", "apaixonado", "apaixonada",
    "focado em resultados", "focada em resultados", "orientado a resultados",
    "orientada a resultados", "fora da caixa", "trabalho em equipe", "multitarefa", "resiliente",
    "comprometido", "comprometida", "perfil inovador", "perfil analítico", "vestir a camisa",
}
FILLER = {
    "in order to": "to", "on a daily basis": "daily", "due to the fact that": "because",
    "a wide variety of": "many", "responsible for": "(lead with an action verb)",
    "tasked with": "(lead with an action verb)", "successfully": "(delete; show it)",
    "various": "(name them or delete)", "helped": "(be specific)", "worked on": "(be specific)",
    "responsável por": "(comece com um verbo de ação)", "responsável pela": "(comece com um verbo de ação)",
    "responsável pelo": "(comece com um verbo de ação)", "encarregado de": "(comece com um verbo de ação)",
    "auxiliei": "(diga exatamente o que você fez)", "ajudei": "(diga exatamente o que você fez)",
    "participei": "(diga qual foi a sua parte)", "com sucesso": "(remova; mostre o resultado)",
    "a fim de": "para", "com o objetivo de": "para", "diversos": "(cite quais ou remova)",
    "diversas": "(cite quais ou remova)", "no dia a dia": "diariamente",
}


def deslop(doc: Doc) -> dict:
    findings: list[dict] = []
    for loc, text in doc.segments:
        low = text.lower()
        for word, repl in INFLATED.items():
            if _has(word, low):
                findings.append({"location": loc, "type": "inflated-verb", "found": word, "suggest": repl})
        for bw in BUZZWORDS:
            if _has(bw, low):
                findings.append({"location": loc, "type": "buzzword", "found": bw,
                                 "suggest": "cut / replace with concrete proof"})
        for f, repl in FILLER.items():
            if _has(f, low):
                findings.append({"location": loc, "type": "filler", "found": f, "suggest": repl})
        if "—" in text or "--" in text:
            findings.append({"location": loc, "type": "punctuation", "found": "em-dash/--",
                             "suggest": "use comma/period/'to'"})
    return {"slop_findings": findings, "clean": not findings}


# ---------------------------------------------------------------------------
# ats_file: hygiene of the actual file a recruiter's ATS will parse
# ---------------------------------------------------------------------------

EMAIL_RE = r"[\w.+-]+@[\w-]+\.[\w.-]+"


def ats_file(path: str | Path, doc: Doc) -> list[dict]:
    """Return [{'level': ok|warn|fail, 'msg': ...}] for the given PDF/DOCX/TXT/JSON file."""
    path = Path(path)
    out: list[dict] = []

    def add(level: str, msg: str) -> None:
        out.append({"level": level, "msg": msg})

    chars = len(doc.text.strip())
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        pages = len(PdfReader(path).pages)
        if chars < 200:
            add("fail", f"almost no selectable text ({chars} chars): scanned/image PDF, ATS will read nothing")
        else:
            add("ok", f"selectable text ({chars} chars)")
        add("ok" if pages <= 2 else "warn", f"{pages} page(s)" + ("" if pages <= 2 else ": aim for 1-2"))
    elif path.suffix.lower() == ".docx":
        from docx import Document

        d = Document(str(path))
        if d.tables:
            add("warn", f"{len(d.tables)} table(s): many ATS scramble table cells; use plain paragraphs")
        if d.inline_shapes:
            add("warn", f"{len(d.inline_shapes)} image(s)/icon(s): invisible to ATS")
        header_text = " ".join(p.text for s in d.sections for p in s.header.paragraphs + s.footer.paragraphs)
        if header_text.strip():
            add("warn", "text in page header/footer: some ATS skip it; keep contact info in the body")
        if not (d.tables or d.inline_shapes or header_text.strip()):
            add("ok", "no tables, images or header/footer text")

    if re.search(EMAIL_RE, doc.text):
        add("ok", "email found")
    else:
        add("fail", "no email address found")
    missing = [s for s in ("experience", "education", "skills") if s not in doc.sections]
    if missing:
        add("warn", f"standard heading not detected: {', '.join(missing)}")
    else:
        add("ok", "standard section headings detected")
    return out
