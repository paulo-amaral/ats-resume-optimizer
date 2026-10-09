"""One normalised view of a resume, built from JSON or from a PDF/DOCX/TXT file.

Every check runs on `Doc`, so people can audit the file they already have without writing JSON.
"""
from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from .i18n import HEADING_SYNONYMS, detect_lang, fold

BULLET_CHARS = "•·▪◦●-*–"


@dataclass
class Doc:
    text: str
    lang: str
    sections: dict[str, str] = field(default_factory=dict)
    bullets: list[str] = field(default_factory=list)
    segments: list[tuple[str, str]] = field(default_factory=list)  # (location, text) for findings


def load_json(path: str | Path) -> dict:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except FileNotFoundError:
        sys.exit(f"error: file not found: {path}")
    except json.JSONDecodeError as e:
        sys.exit(f"error: invalid JSON in {path}: {e}")


def extract_text(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        sys.exit(f"error: file not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        from docx import Document

        d = Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return "\n".join(lines)
    return path.read_text(encoding="utf-8")


def from_resume(resume: dict) -> Doc:
    b = resume.get("basics", {})
    segments: list[tuple[str, str]] = []
    bullets: list[str] = []
    if resume.get("summary"):
        segments.append(("summary", resume["summary"]))
    for i, exp in enumerate(resume.get("experience", [])):
        for j, bl in enumerate(exp.get("bullets", [])):
            segments.append((f"experience[{i}].bullet[{j}]", bl))
            bullets.append(bl)
    for i, proj in enumerate(resume.get("projects", [])):
        for j, bl in enumerate(proj.get("bullets", [])):
            segments.append((f"projects[{i}].bullet[{j}]", bl))
            bullets.append(bl)

    parts = [b.get("headline", ""), resume.get("summary", "")]
    for exp in resume.get("experience", []):
        parts += [exp.get("title", ""), exp.get("company", ""), *exp.get("bullets", [])]
    for edu in resume.get("education", []):
        parts += [edu.get("degree", ""), edu.get("field", ""), edu.get("institution", ""), *edu.get("details", [])]
    skills = [i for g in resume.get("skills", []) for i in g.get("items", [])]
    parts += skills
    for proj in resume.get("projects", []):
        parts += [proj.get("name", ""), proj.get("description", ""), *proj.get("bullets", [])]
    for cert in resume.get("certifications", []):
        parts += [cert.get("name", ""), cert.get("issuer", "")]
    for pub in resume.get("publications", []):
        parts += [pub.get("name", ""), pub.get("venue", "")]

    sections = {
        "summary": resume.get("summary", ""),
        "experience": " ".join(bullets),
        "skills": ", ".join(skills),
        "education": " ".join(e.get("institution", "") for e in resume.get("education", [])),
        "certifications": " ".join(c.get("name", "") for c in resume.get("certifications", [])),
    }
    text = "\n".join(p for p in parts if p)
    return Doc(
        text=text,
        lang=resume.get("lang") or detect_lang(text),
        sections={k: v for k, v in sections.items() if v},
        bullets=bullets,
        segments=segments,
    )


def heading_of(line: str) -> str | None:
    key = fold(line).rstrip(":").strip()
    if not key or len(key) > 40:
        return None
    for canonical, names in HEADING_SYNONYMS.items():
        if key in names:
            return canonical
    return None


def split_sections(text: str) -> dict[str, list[str]]:
    """Group lines under recognised headings. Lines before the first heading go to 'header'."""
    out: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        canonical = heading_of(line)
        if canonical:
            current = canonical
            out.setdefault(current, [])
        else:
            out.setdefault(current, []).append(line)
    return out


def from_text(text: str) -> Doc:
    grouped = split_sections(text)
    segments: list[tuple[str, str]] = []
    bullets: list[str] = []
    line_no = 0
    for raw in text.splitlines():
        line_no += 1
        line = raw.strip()
        if not line:
            continue
        segments.append((f"line {line_no}", line))
        if line[0] in BULLET_CHARS:
            bullets.append(line.lstrip(BULLET_CHARS).strip())
    if not bullets:
        # Some PDF extractors drop bullet glyphs; treat long experience lines as bullets.
        bullets = [ln for ln in grouped.get("experience", []) if len(ln.split()) >= 6]
    sections = {k: " ".join(v) for k, v in grouped.items() if k != "header" and v}
    return Doc(text=text, lang=detect_lang(text), sections=sections, bullets=bullets, segments=segments)


def load_doc(path: str | Path) -> tuple[Doc, dict | None]:
    """Return the Doc plus the raw resume dict when the input is JSON."""
    if Path(path).suffix.lower() == ".json":
        resume = load_json(path)
        return from_resume(resume), resume
    return from_text(extract_text(path)), None


def first_match(pattern: str, text: str) -> str:
    m = re.search(pattern, text, re.IGNORECASE)
    return m.group(0).strip() if m else ""
