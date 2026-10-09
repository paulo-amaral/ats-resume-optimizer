"""Resume JSON -> single-column, ATS-safe .docx or text-based .pdf.

Two steps: `blocks()` turns the resume into a flat list of typed blocks (with localized headings);
each backend only maps block kinds to styles. Templates are style dictionaries, not code.
"""
from __future__ import annotations

from pathlib import Path

from .i18n import HEADINGS

TEMPLATES = {
    "classic": {
        "font_pdf": "Helvetica", "font_docx": "Calibri", "align": "center", "upper": True,
        "body": 10.0, "name": 19.0, "heading": 11.5, "accent": "1a1a1a", "rule": "999999",
        "margin": 0.9, "gap": 1.0,
    },
    "modern": {
        "font_pdf": "Helvetica", "font_docx": "Arial", "align": "left", "upper": False,
        "body": 10.0, "name": 22.0, "heading": 12.5, "accent": "1f4e79", "rule": "1f4e79",
        "margin": 0.85, "gap": 1.1,
    },
    "compact": {
        "font_pdf": "Times-Roman", "font_docx": "Times New Roman", "align": "center", "upper": True,
        "body": 9.5, "name": 16.0, "heading": 10.5, "accent": "000000", "rule": "777777",
        "margin": 0.6, "gap": 0.75,
    },
}

CONTACT_KEYS = ("phone", "email", "location", "linkedin", "github", "website")


def _join(*parts: str, sep: str = " | ") -> str:
    return sep.join(p for p in parts if p)


def _dates(item: dict) -> str:
    return f"{item.get('start', '')} – {item.get('end', '')}".strip(" –")


def blocks(resume: dict) -> list[tuple[str, object]]:
    lang = resume.get("lang", "en")
    h = HEADINGS.get(lang, HEADINGS["en"])
    b = resume.get("basics", {})
    out: list[tuple[str, object]] = [("name", b.get("name", ""))]
    if b.get("headline"):
        out.append(("headline", b["headline"]))
    contact = _join(*(b.get(k, "") for k in CONTACT_KEYS), sep="  |  ")
    if contact:
        out.append(("contact", contact))

    if resume.get("summary"):
        out += [("heading", h["summary"]), ("para", resume["summary"])]

    if resume.get("experience"):
        out.append(("heading", h["experience"]))
        for exp in resume["experience"]:
            out.append(("role", _join(exp.get("title", ""), exp.get("company", ""), sep=" | ")))
            meta = _join(exp.get("location", ""), _dates(exp))
            if meta:
                out.append(("meta", meta))
            out += [("bullet", x) for x in exp.get("bullets", [])]

    if resume.get("projects"):
        out.append(("heading", h["projects"]))
        for proj in resume["projects"]:
            title = proj.get("name", "") + (f" ({proj['link']})" if proj.get("link") else "")
            out.append(("role", title))
            if proj.get("description"):
                out.append(("para", proj["description"]))
            out += [("bullet", x) for x in proj.get("bullets", [])]

    if resume.get("skills"):
        out.append(("heading", h["skills"]))
        out += [("skill", (g.get("category", ""), ", ".join(g.get("items", [])))) for g in resume["skills"]]

    if resume.get("education"):
        out.append(("heading", h["education"]))
        for edu in resume["education"]:
            degree = _join(edu.get("degree", ""), edu.get("field", ""), sep=", ")
            out.append(("role", _join(degree, edu.get("institution", ""), sep=" | ")))
            meta = _join(edu.get("location", ""), _dates(edu))
            if meta:
                out.append(("meta", meta))
            out += [("bullet", x) for x in edu.get("details", [])]

    for key, sub in (("certifications", "issuer"), ("publications", "venue"), ("awards", "issuer")):
        if resume.get(key):
            out.append(("heading", h[key]))
            for item in resume[key]:
                line = _join(item.get("name", ""), _join(item.get(sub, ""), item.get("year", ""), sep=", "),
                             sep=" | ")
                if item.get("link"):
                    line += f" ({item['link']})"
                out.append(("para", line))
    return out


def render(resume: dict, out: str | Path, template: str = "classic", paper: str | None = None) -> Path:
    out = Path(out)
    style = TEMPLATES[template]
    # Brazil and most of the world print on A4; US/Canada on Letter.
    paper = paper or ("a4" if resume.get("lang") == "pt" else "letter")
    if out.suffix.lower() == ".pdf":
        _pdf(blocks(resume), out, style, paper, resume.get("basics", {}).get("name", ""))
    elif out.suffix.lower() == ".docx":
        _docx(blocks(resume), out, style, paper)
    elif out.suffix.lower() == ".tex":
        out.write_text(latex(blocks(resume), style, paper, resume.get("lang", "en")), encoding="utf-8")
    else:
        raise ValueError(f"unknown output format '{out.suffix}' (use .pdf, .docx or .tex)")
    return out


# pdfLaTeX is left out on purpose: it cannot switch off font kerning, and CI showed kerned pairs
# extracting as split words ("T exas", "V arejo"). These engines honour Kerning=Off.
LATEX_ENGINES = ("tectonic", "xelatex", "lualatex")
_TEX_ESCAPES = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
    "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}


def _tex(s: str) -> str:
    return "".join(_TEX_ESCAPES.get(c, c) for c in s)


def latex(items, s: dict, paper: str, lang: str) -> str:
    """Single-column LaTeX source with a correct text layer for ATS extraction.

    Meant for XeLaTeX/LuaLaTeX (kerning off). The pdfLaTeX branch still compiles, with
    glyphtounicode so ligatures and accents map to Unicode, but warns because kerning stays on.
    """
    align = "center" if s["align"] == "center" else "flushleft"
    size = "10pt" if s["body"] < 10.5 else "11pt"
    babel = "brazilian" if lang == "pt" else "english"
    head = [
        "% !TEX program = xelatex",
        "% Compile with XeLaTeX or LuaLaTeX (Overleaf: Menu > Compiler > XeLaTeX) for ATS-safe text.",
        rf"\documentclass[{size},{'a4paper' if paper == 'a4' else 'letterpaper'}]{{article}}",
        r"\usepackage{iftex}",
        r"\ifPDFTeX",
        r"  \usepackage[T1]{fontenc}",
        r"  \usepackage[utf8]{inputenc}",
        r"  \usepackage{lmodern}",
        r"  \input{glyphtounicode}",
        r"  \pdfgentounicode=1",
        r"  \GenericWarning{}{ats-resume: compile with XeLaTeX or LuaLaTeX; pdfLaTeX kerning can split words for ATS}",
        r"\else",
        # Kerned pairs (A-W, V-a) leave gaps that text extractors read as spaces: "A WS", "V arejo".
        r"  \usepackage{fontspec}",
        r"  \defaultfontfeatures{Kerning=Off}",
        # By file name: font lookup by family name fails on some TeX installs (e.g. Tectonic).
        r"  \setmainfont{lmroman10}[Extension=.otf,UprightFont=*-regular,BoldFont=*-bold,"
        r"ItalicFont=*-italic,BoldItalicFont=*-bolditalic]",
        r"\fi",
        rf"\usepackage[{babel}]{{babel}}",
        rf"\usepackage[margin={s['margin']}in]{{geometry}}",
        r"\usepackage[hidelinks,pdfusetitle]{hyperref}",
        r"\usepackage{xcolor}",
        r"\usepackage{enumitem}",
        r"\usepackage{titlesec}",
        rf"\definecolor{{accent}}{{HTML}}{{{s['accent'].upper()}}}",
        r"\pagestyle{empty}",
        r"\setlength{\parindent}{0pt}",
        r"\setlist[itemize]{leftmargin=1.2em,itemsep=1pt,topsep=1pt,parsep=0pt}",
        r"\titleformat{\section}{\large\bfseries\color{accent}" + (r"\MakeUppercase" if s["upper"] else "")
        + r"}{}{0pt}{}[\vspace{-0.6em}\rule{\linewidth}{0.4pt}]",
        r"\titlespacing*{\section}{0pt}{0.8em}{0.4em}",
        r"\begin{document}",
    ]
    body: list[str] = [rf"\begin{{{align}}}"]
    in_list = False
    in_header = True  # name/headline/contact share one aligned block, closed before the first section
    for kind, value in items:
        if kind == "bullet":
            if not in_list:
                body.append(r"\begin{itemize}")
                in_list = True
            body.append(rf"  \item {_tex(value)}")
            continue
        if in_list:
            body.append(r"\end{itemize}")
            in_list = False
        if in_header and kind not in ("name", "headline", "contact"):
            body.append(rf"\end{{{align}}}")
            in_header = False
        if kind == "name":
            body.append(rf"{{\LARGE\bfseries\color{{accent}} {_tex(value)}}}\\[2pt]")
        elif kind == "headline":
            body.append(rf"{{\large {_tex(value)}}}\\[2pt]")
        elif kind == "contact":
            body.append(rf"{{\small {_tex(value)}}}")
        elif kind == "heading":
            body.append(rf"\section*{{{_tex(value)}}}")
        elif kind == "role":
            body.append(rf"\textbf{{{_tex(value)}}}\\")
        elif kind == "meta":
            body.append(rf"{{\small\itshape {_tex(value)}}}\par")
        elif kind == "skill":
            body.append(rf"\textbf{{{_tex(value[0])}:}} {_tex(value[1])}\par")
        else:
            body.append(rf"{_tex(value)}\par")
    if in_list:
        body.append(r"\end{itemize}")
    if in_header:
        body.append(rf"\end{{{align}}}")
    return "\n".join(head + body + [r"\end{document}", ""])


def compile_latex(tex: Path) -> Path:
    """Compile .tex -> .pdf with the first LaTeX engine found on PATH."""
    import shutil
    import subprocess

    engine = next((e for e in LATEX_ENGINES if shutil.which(e)), None)
    if engine is None:
        raise RuntimeError(
            "no XeLaTeX/LuaLaTeX engine found (install tectonic, or TeX Live / MiKTeX); "
            f"the .tex file is ready: {tex}  (you can also upload it to Overleaf)"
        )
    cmd = [engine, tex.name] if engine == "tectonic" else [engine, "-interaction=nonstopmode", "-halt-on-error", tex.name]
    result = subprocess.run(cmd, cwd=tex.parent, capture_output=True, text=True)
    if result.returncode != 0:
        log = (result.stdout + result.stderr).strip().splitlines()
        errors = [ln for ln in log if ln.startswith(("!", "error")) or "Error" in ln] or log[-5:]
        raise RuntimeError(f"{engine} failed on {tex}:\n  " + "\n  ".join(errors[:6]))
    return tex.with_suffix(".pdf")


def _docx(items, out: Path, s: dict, paper: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        if paper == "a4":
            section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.left_margin = section.right_margin = Inches(s["margin"] + 0.1)
        section.top_margin = section.bottom_margin = Inches(s["margin"])
    normal = doc.styles["Normal"]
    normal.font.name = s["font_docx"]
    normal.font.size = Pt(s["body"] + 0.5)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), s["font_docx"])
    rfonts.set(qn("w:hAnsi"), s["font_docx"])
    accent = RGBColor.from_string(s["accent"].upper())
    align = WD_ALIGN_PARAGRAPH.CENTER if s["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT

    def para(space_after: float = 2.0, style_name: str | None = None):
        p = doc.add_paragraph(style=style_name)
        p.paragraph_format.space_after = Pt(space_after * s["gap"])
        return p

    for kind, value in items:
        if kind in ("name", "headline", "contact"):
            p = para(0 if kind != "contact" else 4)
            p.alignment = align
            r = p.add_run(value)
            r.font.size = Pt({"name": s["name"], "headline": s["heading"], "contact": s["body"] - 0.5}[kind])
            r.bold = kind == "name"
            if kind == "name":
                r.font.color.rgb = accent
        elif kind == "heading":
            p = para()
            p.paragraph_format.space_before = Pt(10 * s["gap"])
            r = p.add_run(value.upper() if s["upper"] else value)
            r.bold = True
            r.font.size = Pt(s["heading"])
            r.font.color.rgb = accent
            ppr = p._p.get_or_add_pPr()
            pbdr = ppr.makeelement(qn("w:pBdr"), {})
            pbdr.append(pbdr.makeelement(qn("w:bottom"), {
                qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): s["rule"],
            }))
            ppr.append(pbdr)
        elif kind == "role":
            para(0).add_run(value).bold = True
        elif kind == "meta":
            r = para().add_run(value)
            r.italic = True
            r.font.size = Pt(s["body"] - 0.5)
        elif kind == "bullet":
            para(style_name="List Bullet").add_run(value)
        elif kind == "skill":
            p = para()
            p.add_run(f"{value[0]}: ").bold = True
            p.add_run(value[1])
        else:
            para().add_run(value)
    doc.save(str(out))


def _pdf(items, out: Path, s: dict, paper: str, author: str) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate

    font = s["font_pdf"]
    bold = {"Helvetica": "Helvetica-Bold", "Times-Roman": "Times-Bold"}[font]
    italic = {"Helvetica": "Helvetica-Oblique", "Times-Roman": "Times-Italic"}[font]
    align = TA_CENTER if s["align"] == "center" else TA_LEFT
    g = s["gap"]
    accent = HexColor("#" + s["accent"])
    body = ParagraphStyle("body", fontName=font, fontSize=s["body"], leading=s["body"] * 1.25, spaceAfter=3 * g)
    st = {
        "name": ParagraphStyle("name", parent=body, fontName=bold, fontSize=s["name"], leading=s["name"] * 1.15,
                               alignment=align, textColor=accent, spaceAfter=1),
        "headline": ParagraphStyle("headline", parent=body, fontSize=s["body"] + 1, alignment=align, spaceAfter=1),
        "contact": ParagraphStyle("contact", parent=body, fontSize=s["body"] - 1, alignment=align,
                                  textColor=HexColor("#333333"), spaceAfter=6 * g),
        "heading": ParagraphStyle("heading", parent=body, fontName=bold, fontSize=s["heading"],
                                  spaceBefore=6 * g, spaceAfter=2, textColor=accent),
        "role": ParagraphStyle("role", parent=body, fontName=bold, fontSize=s["body"] + 0.5, spaceBefore=2 * g,
                               spaceAfter=0),
        "meta": ParagraphStyle("meta", parent=body, fontName=italic, fontSize=s["body"] - 1,
                               textColor=HexColor("#444444"), spaceAfter=2 * g),
        "bullet": ParagraphStyle("bullet", parent=body, spaceAfter=1.5 * g),
    }

    flow: list = []
    pending: list = []

    def flush() -> None:
        if pending:
            flow.append(ListFlowable([ListItem(p, leftIndent=10, value="•") for p in pending],
                                     bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8))
            pending.clear()

    for kind, value in items:
        if kind == "bullet":
            pending.append(Paragraph(escape(value), st["bullet"]))
            continue
        flush()
        if kind == "heading":
            flow.append(Paragraph(escape(value.upper() if s["upper"] else value), st["heading"]))
            flow.append(HRFlowable(width="100%", thickness=0.6, color=HexColor("#" + s["rule"]),
                                   spaceBefore=1, spaceAfter=4))
        elif kind == "skill":
            flow.append(Paragraph(f"<b>{escape(value[0])}:</b> {escape(value[1])}", body))
        else:
            flow.append(Paragraph(escape(value), st.get(kind, body)))
    flush()

    m = s["margin"] * inch
    SimpleDocTemplate(
        str(out), pagesize=A4 if paper == "a4" else LETTER, leftMargin=m, rightMargin=m,
        topMargin=m * 0.7, bottomMargin=m * 0.6, title=f"{author} - Resume", author=author,
    ).build(flow)
